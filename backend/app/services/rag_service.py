import os
import uuid
import pypdf
import docx
import numpy as np
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_google_genai import GoogleGenerativeAIEmbeddings
from sqlalchemy.orm import Session
from ..database import models
from ..config import settings
from typing import List

def extract_text_from_file(file_path: str, file_type: str) -> str:
    text = ""
    try:
        if file_type == "application/pdf" or file_path.lower().endswith(".pdf"):
            reader = pypdf.PdfReader(file_path)
            for page in reader.pages:
                text += (page.extract_text() or "") + "\n"
        elif file_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" or file_path.lower().endswith(".docx"):
            doc = docx.Document(file_path)
            for para in doc.paragraphs:
                text += para.text + "\n"
        else:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
    except Exception as e:
        print(f"Error extracting text from file {file_path}: {e}")
    return text

def process_document(file_path: str, filename: str, file_type: str, conversation_id: int, db: Session):
    text = extract_text_from_file(file_path, file_type)
    
    # Save document record
    doc_record = models.Document(
        conversation_id=conversation_id,
        filename=filename,
        file_type=file_type,
        file_path=file_path,
        extracted_text=text
    )
    db.add(doc_record)
    db.commit()
    db.refresh(doc_record)
    
    # Chunk text
    text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200)
    chunks = text_splitter.split_text(text) if text else []
    if not chunks and text:
        chunks = [text]
    
    # Generate embeddings if API key is present
    embeddings = None
    api_key = settings.LLM_API_KEY or os.environ.get("GEMINI_API_KEY") or os.environ.get("GOOGLE_API_KEY")
    if api_key and chunks:
        try:
            embeddings_model = GoogleGenerativeAIEmbeddings(model="models/embedding-001", google_api_key=api_key)
            embeddings = embeddings_model.embed_documents(chunks)
        except Exception as e:
            print(f"Embeddings generation error: {e}")
            embeddings = None
        
    # Save chunks
    for i, chunk in enumerate(chunks):
        chunk_record = models.DocumentChunk(
            document_id=doc_record.id,
            chunk_text=chunk,
            embedding=embeddings[i] if embeddings and i < len(embeddings) else None
        )
        db.add(chunk_record)
    db.commit()
        
    return doc_record.id

def search_similar_chunks(query: str, conversation_id: int, db: Session, limit: int = 3) -> List[str]:
    api_key = settings.LLM_API_KEY or os.environ.get("GEMINI_API_KEY") or os.environ.get("GOOGLE_API_KEY")
    
    # Try embedding search if API key exists
    if api_key:
        try:
            embeddings_model = GoogleGenerativeAIEmbeddings(model="models/embedding-001", google_api_key=api_key)
            query_embedding = embeddings_model.embed_query(query)
            
            from ..database.database import is_sqlite
            if not is_sqlite:
                try:
                    results = db.query(models.DocumentChunk).join(models.Document).filter(
                        models.Document.conversation_id == conversation_id
                    ).order_by(models.DocumentChunk.embedding.l2_distance(query_embedding)).limit(limit).all()
                    if results:
                        return [res.chunk_text for res in results if res.chunk_text]
                except Exception as e:
                    print(f"pgvector query error: {e}")
            
            # Python cosine similarity fallback
            chunks = db.query(models.DocumentChunk).join(models.Document).filter(
                models.Document.conversation_id == conversation_id
            ).all()
            scored = []
            q_arr = np.array(query_embedding, dtype=float)
            q_norm = np.linalg.norm(q_arr)
            for c in chunks:
                if c.embedding:
                    c_arr = np.array(c.embedding, dtype=float)
                    c_norm = np.linalg.norm(c_arr)
                    if q_norm > 0 and c_norm > 0:
                        sim = float(np.dot(q_arr, c_arr) / (q_norm * c_norm))
                        scored.append((sim, c.chunk_text))
            scored.sort(key=lambda x: x[0], reverse=True)
            if scored:
                return [item[1] for item in scored[:limit]]
        except Exception as e:
            print(f"Vector search failed: {e}")

    # Fallback keyword matching when embeddings aren't available or no chunks found
    query_words = set(w.lower() for w in query.split() if len(w) > 2)
    chunks = db.query(models.DocumentChunk).join(models.Document).filter(
        models.Document.conversation_id == conversation_id
    ).all()
    scored = []
    for c in chunks:
        score = sum(1 for w in query_words if w in c.chunk_text.lower())
        if score > 0:
            scored.append((score, c.chunk_text))
    scored.sort(key=lambda x: x[0], reverse=True)
    if scored:
        return [item[1] for item in scored[:limit]]
    elif chunks:
        return [c.chunk_text for c in chunks[:limit]]
    return []

